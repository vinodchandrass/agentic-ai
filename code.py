import openpyxl
import pandas as pd

excel_path = "ArticleC_Native_Reviewer_Package_180Items_v1.xlsx"
wb = openpyxl.load_workbook(excel_path)
print("Sheet names:", wb.sheetnames)

for name in wb.sheetnames:
    df = pd.read_excel(excel_path, sheet_name=name)
    print(f"\n--- Sheet: {name} ---")
    print(df.head(2))
    print("Columns:", df.columns.tolist())